"""Independent workbook mirrors; synthetic business messages use actual processing entries."""
import unittest,tempfile,os,asyncio,json,hashlib
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch,AsyncMock
from openpyxl import Workbook,load_workbook
from team_bot.desktop_sync import DesktopSync,FILES,destination
import test_progress_rounds as fixtures

class Settings:
 def __init__(self):self.values={}
 def setting(self,key,default=None):return self.values.get(key,default)

def book(path,names,value):
 path.parent.mkdir(parents=True,exist_ok=True);w=Workbook();w.remove(w.active)
 for name in names:w.create_sheet(name)['A1']=value
 w.save(path);w.close()

class MirrorTests(unittest.IsolatedAsyncioTestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.target=self.root/'existing desktop folder';self.target.mkdir()
  self.s=SimpleNamespace(c=SimpleNamespace(data=self.root),db=Settings());self.p=SimpleNamespace(c=self.s.c,db=Settings())
  self.daily=self.root/'exports'/FILES['daily'][0];self.progress=self.root/'exports'/FILES['progress'][0]
  book(self.daily,FILES['daily'][1],1);book(self.progress,FILES['progress'][1],2)
  self.s.db.values.update(daily_excel_path=str(self.daily),daily_excel_revision=1,daily_excel_done=1,daily_excel_pending=False)
  self.p.db.values.update(progress_excel_path=str(self.progress),progress_excel_revision=1,progress_excel_done=1,progress_excel_dirty=False)
  self.env=patch.dict(os.environ,{'PROGRESS_EXCEL_PATH':''});self.env.start();self.sync=DesktopSync(self.s,str(self.target),self.p)
 def tearDown(self):self.env.stop();self.tmp.cleanup()
 async def test_both_fixed_names_byte_identical_and_cache_state_stays_synced(self):
  for _ in range(3):await self.sync.once()
  for kind,src in [('daily',self.daily),('progress',self.progress)]:
   self.assertEqual(src.read_bytes(),(self.target/src.name).read_bytes());self.assertFalse(self.sync.state[kind]['pending'])
  self.assertEqual(sorted(x.name for x in self.target.iterdir()),sorted([self.daily.name,self.progress.name]))
 async def test_source_pending_and_invalid_never_replace_valid_copy(self):
  await self.sync.once();old=(self.target/self.daily.name).read_bytes();book(self.daily,FILES['daily'][1],7)
  self.s.db.values['daily_excel_pending']=True;await self.sync.once();self.assertEqual(old,(self.target/self.daily.name).read_bytes())
  self.s.db.values['daily_excel_pending']=False;self.daily.write_bytes(b'invalid');await self.sync.once();self.assertEqual(old,(self.target/self.daily.name).read_bytes());self.assertTrue(self.sync.state['daily']['pending'])
 async def test_restart_reconciles_modified_or_deleted_target(self):
  await self.sync.once();(self.target/self.daily.name).write_bytes(b'wrong');(self.target/self.progress.name).unlink()
  restarted=DesktopSync(self.s,str(self.target),self.p);await restarted.once()
  self.assertEqual(self.daily.read_bytes(),(self.target/self.daily.name).read_bytes());self.assertEqual(self.progress.read_bytes(),(self.target/self.progress.name).read_bytes())
 async def test_absent_folder_not_created_and_configuration_retry(self):
  missing=self.root/'not created';other=DesktopSync(self.s,str(missing),self.p);await other.once();self.assertFalse(missing.exists());self.assertTrue(other.state['daily']['pending'])
  missing.mkdir();await other.once();self.assertFalse(other.state['daily']['pending'])
 async def test_windows_actual_file_lock_retries_and_keeps_previous(self):
  if os.name!='nt':self.skipTest('Windows sharing lock')
  await self.sync.once();dest=self.target/self.daily.name;old=dest.read_bytes();book(self.daily,FILES['daily'][1],8)
  import ctypes
  from ctypes import wintypes
  k=ctypes.WinDLL('kernel32',use_last_error=True);k.CreateFileW.argtypes=[wintypes.LPCWSTR,wintypes.DWORD,wintypes.DWORD,ctypes.c_void_p,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE];k.CreateFileW.restype=wintypes.HANDLE;k.CloseHandle.argtypes=[wintypes.HANDLE]
  handle=k.CreateFileW(str(dest),0x80000000,1,None,3,0,None);self.assertNotEqual(handle,ctypes.c_void_p(-1).value)
  try:
   await self.sync.once();self.assertEqual(dest.read_bytes(),old);self.assertTrue(self.sync.state['daily']['pending']);self.assertEqual(self.sync.state['daily']['error'],'PermissionError')
  finally:k.CloseHandle(handle)
  await DesktopSync(self.s,str(self.target),self.p).once();self.assertEqual(dest.read_bytes(),self.daily.read_bytes())
 async def test_one_failure_does_not_block_other_workbook(self):
  await self.sync.once();book(self.daily,FILES['daily'][1],9);book(self.progress,FILES['progress'][1],9)
  real=os.replace
  def replace(a,b):
   if Path(b)==self.target/self.daily.name:raise PermissionError('mock occupied')
   return real(a,b)
  with patch('team_bot.desktop_sync.os.replace',side_effect=replace):await self.sync.once()
  self.assertTrue(self.sync.state['daily']['pending']);self.assertFalse(self.sync.state['progress']['pending']);self.assertEqual((self.target/self.progress.name).read_bytes(),self.progress.read_bytes())
 def test_desktop_resolution_uses_system_and_config_default_disabled(self):
  self.assertIsNone(destination(''))
  if os.name=='nt':self.assertTrue(destination('desktop:工作历史记录表').parent.is_dir())

class BusinessEntryTests(unittest.IsolatedAsyncioTestCase):
 async def test_saved_daily_query_and_miniapp_feedback_sync_separately(self):
  f=fixtures.ProgressTests();f.setUp();target=f.c.data/'mirror';target.mkdir();f.p.db.set('progress_excel_path',str(f.c.data/'exports'/FILES['progress'][0]))
  try:
   from team_bot.daily_history import ensure_export,export_path
   from team_bot.progress_export import ensure_export as ensure_progress,export_path as progress_path
   from team_bot.progress_miniapp import MiniAPI
   from test_progress_miniapp import signed
   # Mirror logic must not depend on the host's optional proprietary renderer.
   # The deployed compatibility writers still generate real, validated XLSX files.
   from team_bot.portable_excel import write_daily
   writer=patch('team_bot.daily_history.write_excel',new=write_daily);writer.start()
   env=patch.dict(os.environ,{'PROGRESS_EXCEL_PATH':'','PROGRESS_MINIAPP_LINK':'https://t.me/mock_company_bot/progress','DAILY_ARTIFACT_MODULE':str(f.c.data/'no-artifact-provider')});env.start()
   sync=DesktopSync(f.s,str(target),f.p)
   await f.f.handle(710,f.u,'2026-10-10\n新增：2\n维聊：3\n有效：4\n精聊：0\n备注：模拟同步。')
   self.assertTrue(await ensure_export(f.s,True));self.assertTrue(await ensure_progress(f.p));await sync.once()
   daily=(target/FILES['daily'][0]).read_bytes();self.assertEqual(daily,export_path(f.s).read_bytes())
   # New query uses the actual Mini App API, not a direct record insertion.
   api=MiniAPI(f.p,f.ctx.bot);api.membership=AsyncMock()
   q=await api.dispatch('/api/query',dict(initData=signed(f.c.token,f.c.owner),body=dict(day='2026-10-10',key='mock-desktop-query')))
   rid=q['round'];self.assertTrue(await ensure_progress(f.p));await sync.once();query_bytes=(target/FILES['progress'][0]).read_bytes()
   await api.dispatch('/api/submit',dict(initData=signed(f.c.token,f.u),body=dict(round=rid,key='mock-desktop-feedback',masks={str(f.t):3,str(f.z):5},notes={str(f.t):'模拟当前问题',str(f.z):''})))
   self.assertTrue(await ensure_progress(f.p));await sync.once();self.assertNotEqual(query_bytes,(target/FILES['progress'][0]).read_bytes());self.assertEqual((target/FILES['progress'][0]).read_bytes(),progress_path(f.p).read_bytes());self.assertEqual(daily,(target/FILES['daily'][0]).read_bytes())
   await f.f.handle(711,f.u,'2026-10-10\n新增：5\n维聊：3\n有效：4\n精聊：0')
   self.assertTrue(await ensure_export(f.s,True));await sync.once();self.assertNotEqual(daily,(target/FILES['daily'][0]).read_bytes());self.assertEqual((target/FILES['daily'][0]).read_bytes(),export_path(f.s).read_bytes())
  finally:
   if 'env' in locals():env.stop()
   if 'writer' in locals():writer.stop()
   f.tearDown()

if __name__=='__main__':unittest.main()
