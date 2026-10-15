import unittest,tempfile,os,json
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from team_bot.progress_export import write_export
from team_bot.progress_portable_xlsx import write

class ProgressExcelTests(unittest.TestCase):
 def data(self):
  note='独立模拟完整多行问题。'*150+'\n第二行不会丢失。'
  task=dict(id=1,version=1,title='独立模拟任务',kind='free',stages=['一','二','三','四'])
  rows=[['2026-10-10','13:00','模拟员工'+str(i),'mock'+str(i),'已提交',1,0,1,0,2,4,'2/4',2,note if i==0 else '', '10-10 13:00',1,1,str(555100+i),'进行中'] for i in range(30)]
  return dict(label='模拟展示，不计正式数据',revision=1,summaries=[['2026-10-10','13:00','独立模拟任务',1,30,30,0,0,1,1,'',1]],groups=[dict(task=task,rows=rows)],problems=[['2026-10-10','13:00','模拟员工0','mock0','独立模拟任务','2/4',2,note,'10-10 13:00',1,1,'555100']],definitions=[[1,'2026-10-10','独立模拟任务','自由完成',4,'一','二','三','四','']]),note
 def inspect(self,path,note):
  wb=load_workbook(path);self.assertEqual(wb.sheetnames,['查询汇总','员工进度历史','当前问题历史','任务定义'])
  sh=wb['员工进度历史'];self.assertIn('A5:S5',[str(x) for x in sh.merged_cells]);self.assertEqual(sh.max_row,36)
  self.assertEqual(sh['N7'].value,note)
  issues=wb['当前问题历史'];self.assertEqual(issues.max_row,6);self.assertEqual(issues['H6'].value,note)
  self.assertEqual(issues['C6'].value,'模拟员工0');self.assertEqual(issues['K6'].value,1)
  self.assertEqual(sh['A7'].number_format,'yyyy-mm-dd');self.assertEqual(sh['F7'].value,1);self.assertIsInstance(sh['F7'].value,int)
  for row in sh.iter_rows(min_row=7):
   for cell in row:
    if cell.value is not None:
     self.assertEqual(cell.alignment.horizontal,'center');self.assertTrue(cell.alignment.wrap_text);self.assertEqual(cell.border.left.style,'thin')
  self.assertTrue(all(dim.height is None or dim.height<=180 for dim in sh.row_dimensions.values()))
 def test_portable_complete_style_long_notes_and_expand(self):
  d,note=self.data()
  with tempfile.TemporaryDirectory() as tmp:
   dest=Path(tmp)/'v2.xlsx';write(d,dest);self.inspect(dest,note)
 def test_atomic_busy_file_preserves_previous_and_retry_uses_same_snapshot(self):
  d,note=self.data()
  with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,DAILY_ARTIFACT_MODULE=tmp):
   dest=Path(tmp)/'v2.xlsx';dest.write_bytes(b'previous-valid-placeholder')
   with patch('team_bot.progress_export.os.replace',side_effect=PermissionError('模拟打开占用')):
    with self.assertRaises(ValueError):write_export(d,dest)
   self.assertEqual(dest.read_bytes(),b'previous-valid-placeholder');write_export(d,dest);self.inspect(dest,note)
 def test_installed_generator_one_row_issue_literal_and_anchor_only(self):
  import zipfile,xml.etree.ElementTree as E
  d,note=self.data();d['problems'][0][7]='=原文而非公式\n'+note;d['groups'][0]['rows'][0][13]=d['problems'][0][7]
  with tempfile.TemporaryDirectory() as tmp:
   dest=Path(tmp)/'actual-generator.xlsx';write_export(d,dest);self.inspect(dest,d['problems'][0][7]);w=load_workbook(dest);self.assertEqual(w['当前问题历史']['H6'].data_type,'s')
   with zipfile.ZipFile(dest) as z:
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    for n in range(1,5):
     root=E.fromstring(z.read(f'xl/worksheets/sheet{n}.xml'));cells={c.get('r'):c for c in root.findall('.//s:sheetData/s:row/s:c',ns)}
     for merge in root.findall('s:mergeCells/s:mergeCell',ns):
      from openpyxl.utils.cell import range_boundaries
      a,b,c,e=range_boundaries(merge.get('ref'))
      for row in range(b,e+1):
       for col in range(a,c+1):
        from openpyxl.utils import get_column_letter
        key=get_column_letter(col)+str(row)
        if row==b and col==a:continue
        cell=cells.get(key)
        if cell is not None:self.assertFalse(cell.find('s:v',ns) is not None and cell.find('s:v',ns).text,'Non-anchor merged title contains a duplicate')
 def test_issue_text_matching_status_is_not_restyled_as_task_completion(self):
  d,_=self.data();d['problems'][0][7]='已完成';d['groups'][0]['rows'][0][13]='已完成';d['groups'][0]['rows'][0][5]='已完成'
  with tempfile.TemporaryDirectory() as tmp:
   for writer in (write,write_export):
    dest=Path(tmp)/'style.xlsx';writer(d,dest);book=load_workbook(dest)
    self.assertEqual(book['当前问题历史']['H6'].value,'已完成');self.assertNotEqual(book['当前问题历史']['H6'].fill.fgColor.rgb[-6:],'E2F0D9');self.assertEqual(book['员工进度历史']['F7'].fill.fgColor.rgb[-6:],'E2F0D9')
