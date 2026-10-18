"""Real no-artifact/no-Node Excel output, against existing provider contracts."""
from copy import deepcopy
from datetime import date, datetime
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from team_bot import config, daily_history, portable_excel, progress_export
from team_bot.progress_portable_xlsx import write as portable_progress
import test_progress_excel_v2 as progress_fixtures


def daily_fixture():
    days=[]
    for day in ('2026-10-09','2026-10-10','2026-10-11'):
        rows=[['模拟员工'+str(i),'@mock'+str(i),i,0,2,1,
               ('=完整问题原文，不能当公式。\n第二行。'*30) if i==0 else '', '已提交 · 迟交'] for i in range(30)]
        days.append(dict(day=day,submitted=30,expected=30,status='已齐',
                         totals=dict(新增=435,维聊=0,有效=60,精聊=30),rows=rows))
    return dict(days=days,label='独立模拟，不计正式数据',revision=3)


def workbook_contract(path):
    """Compare semantic content/style, not volatile ZIP metadata or style IDs."""
    book=load_workbook(path)
    try:
        out={}
        for sh in book:
            cells={}
            for row in sh:
                for c in row:
                    if c.value is None: continue
                    value=c.value
                    if isinstance(value,datetime) and value.time().isoformat()=='00:00:00': value=value.date()
                    cells[c.coordinate]=(value,c.number_format,c.alignment.horizontal,c.alignment.vertical,
                                         c.alignment.wrap_text,c.font.name,c.font.sz,c.font.bold,
                                         c.fill.fgColor.type,c.fill.fgColor.rgb if c.fill.fgColor.type=='rgb' else c.fill.fgColor.indexed,
                                         c.border.left.style,c.border.right.style,c.data_type)
            out[sh.title]=dict(cells=cells,merged=sorted(str(r) for r in sh.merged_cells),freeze=sh.freeze_panes,
                               widths={k:v.width for k,v in sh.column_dimensions.items()},
                               heights={k:v.height for k,v in sh.row_dimensions.items()})
        return out
    finally: book.close()


class Stage3ExportTests(unittest.TestCase):
    def assert_public_format(self, daily_path, progress_path, daily, progress):
        """Independent input/template assertions, not writer-vs-itself comparison."""
        from openpyxl.utils import get_column_letter
        book=load_workbook(daily_path)
        try:
            self.assertEqual(book.sheetnames,['每日汇总','员工明细'])
            summary,detail=book.worksheets
            self.assertEqual([c.value for c in summary[4]],['日期','新增','维聊','有效','精聊','已交人数','应交人数','状态'])
            self.assertEqual([summary.cell(5,i).value for i in range(2,8)],[435,0,60,30,30,30])
            self.assertEqual(summary['A5'].value.date().isoformat(),daily['days'][0]['day'])
            for row in summary.iter_rows(min_row=4,max_row=7):
                for c in row:
                    self.assertEqual((c.alignment.horizontal,c.alignment.vertical),('center','center'))
            self.assertEqual(detail.freeze_panes,'A3')
            for row in (4,):
                self.assertIn(f'A{row}:H{row}',[str(x) for x in detail.merged_cells])
                self.assertEqual(detail.cell(row,1).number_format,'yyyy-mm-dd')
            for col in (1,2,7):self.assertEqual(detail.cell(6,col).alignment.horizontal,'left')
            for col in (3,4,5,6):
                self.assertEqual(detail.cell(6,col).alignment.horizontal,'center')
                self.assertEqual(detail.column_dimensions[get_column_letter(col)].width,11)
                self.assertIsInstance(detail.cell(6,col).value,int)
            self.assertTrue(detail['G6'].alignment.wrap_text)
            notes=''.join(row[6].value or '' for row in detail.iter_rows(min_row=6) if row[6].value and row[6].value!='备注')
            self.assertEqual(notes,''.join(str(r[6]) for d in daily['days'] for r in d['rows']))
        finally: book.close()
        book=load_workbook(progress_path)
        template=load_workbook(config.ROOT/'templates/Telegram_Progress_History_Template_v2.xlsx')
        try:
            self.assertEqual(book.sheetnames,['查询汇总','员工进度历史','当前问题历史','任务定义'])
            for name,last,head in [('查询汇总',12,5),('员工进度历史',19,6),('当前问题历史',12,5),('任务定义',10,5)]:
                sh,ref=book[name],template[name]
                self.assertEqual(sh.freeze_panes,'A6')
                self.assertEqual([sh.cell(head,i).value for i in range(1,last+1)],
                    [str(ref.cell(head,i).value).replace('身份ID（演示）','身份ID') for i in range(1,last+1)])
                for i in range(1,last+1):
                    col=get_column_letter(i)
                    self.assertAlmostEqual(sh.column_dimensions[col].width,ref.column_dimensions[col].width,delta=.2)
                    c=sh.cell(head+1,i)
                    if c.value is not None:
                        self.assertEqual((c.alignment.horizontal,c.alignment.vertical),('center','center'))
                        self.assertTrue(c.alignment.wrap_text)
                        for edge in ('left','right','top','bottom'):self.assertEqual(getattr(c.border,edge).style,'thin')
                for r in sh.merged_cells:
                    for row in sh.iter_rows(min_row=r.min_row,max_row=r.max_row,min_col=r.min_col,max_col=r.max_col):
                        for c in row:
                            if c.coordinate!=r.start_cell.coordinate:self.assertIsNone(c.value)
            self.assertEqual(book['当前问题历史'].max_row,5+len(progress['problems']))
            for i,original in enumerate(progress['problems'],6):
                self.assertEqual(book['当前问题历史'].cell(i,8).value,original[7])
            self.assertEqual(book['员工进度历史']['F7'].fill.fgColor.rgb[-6:],'E2F0D9')
            self.assertEqual(book['员工进度历史']['N7'].value,progress['groups'][0]['rows'][0][13])
            self.assertIn('A38:S38',[str(x) for x in book['员工进度历史'].merged_cells])
            self.assertEqual(book['员工进度历史']['A7'].number_format,'yyyy-mm-dd')
        finally: book.close();template.close()

    def test_reproducible_public_format_contract_both_real_providers(self):
        # Optional provider absent => explicit skip for that subtest, never a
        # fake module import counted as successful artifact generation.
        daily=daily_fixture();progress,_=progress_fixtures.ProgressExcelTests().data()
        progress['groups'][0]['rows'][0][5]='已完成'
        second=deepcopy(progress['groups'][0]);second['task'].update(id=2,title='第二项独立模拟任务')
        progress['groups'].append(second)
        before=deepcopy((daily,progress))
        runtime=portable_excel.artifact_runtime()
        for provider in ('openpyxl','artifact'):
            with self.subTest(provider=provider),tempfile.TemporaryDirectory(prefix='excel-contract-') as td:
                folder=Path(td)
                if provider=='artifact' and runtime is None:
                    self.skipTest('可选Node/artifact运行时未安装；仅openpyxl路径已实测')
                with patch.object(portable_excel,'artifact_runtime',return_value=runtime if provider=='artifact' else None):
                    daily_history.write_excel(daily,folder/'daily.xlsx')
                    progress_export.write_export(progress,folder/'progress.xlsx')
                self.assert_public_format(folder/'daily.xlsx',folder/'progress.xlsx',daily,progress)
                self.assertEqual((daily,progress),before)

    def test_real_two_exports_without_codex_node_or_artifact(self):
        daily=daily_fixture();progress,note=progress_fixtures.ProgressExcelTests().data()
        before=deepcopy((daily,progress))
        with tempfile.TemporaryDirectory(prefix='stage3-portable-') as td:
            root=Path(td);(root/'templates').mkdir()
            import shutil
            shutil.copyfile(config.ROOT/'templates/Telegram_Progress_History_Template_v2.xlsx',root/'templates/Telegram_Progress_History_Template_v2.xlsx')
            target=root/'exports';target.mkdir()
            with patch.object(config,'ROOT',root),patch.object(Path,'home',return_value=root/'clean-user'), \
                 patch.object(portable_excel.shutil,'which',return_value=None), \
                 patch.dict(os.environ,DAILY_NODE='',DAILY_ARTIFACT_MODULE=''), \
                 patch('subprocess.run',side_effect=AssertionError('portable export must not spawn Node')):
                self.assertFalse(portable_excel.available_artifact())
                portable_excel.install_if_missing()
                daily_history.write_excel(daily,target/'日报历史记录.xlsx')
                progress_export.write_export(progress,target/'工作进度历史.xlsx')
                portable_excel.write_daily(daily,target/'daily-reference.xlsx')
                portable_progress(progress,target/'progress-reference.xlsx')
            self.assertEqual(workbook_contract(target/'日报历史记录.xlsx'),workbook_contract(target/'daily-reference.xlsx'))
            self.assertEqual(workbook_contract(target/'工作进度历史.xlsx'),workbook_contract(target/'progress-reference.xlsx'))
            book=load_workbook(target/'日报历史记录.xlsx')
            self.assertEqual(book.sheetnames,['每日汇总','员工明细'])
            self.assertEqual(book['每日汇总']['A5'].number_format,'yyyy-mm-dd')
            self.assertEqual(book['每日汇总']['B5'].value,435)
            self.assertEqual(book['每日汇总']['C5'].value,0)
            self.assertEqual(book['员工明细'].freeze_panes,'A3')
            # Reassemble established labelled daily continuation rows; no text loss.
            notes=''.join(c.value or '' for row in book['员工明细'].iter_rows(min_row=6) for c in [row[6]] if c.value)
            self.assertEqual(notes.count('第二行。'),90)
            self.assertTrue(all(c.data_type=='s' for row in book['员工明细'] for c in [row[6]] if isinstance(c.value,str)))
            book.close()
            progress_fixtures.ProgressExcelTests().inspect(target/'工作进度历史.xlsx',note)
            self.assertEqual((daily,progress),before)

    def test_existing_optional_artifact_is_still_selected(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);(root/'node_modules/@oai/artifact-tool').mkdir(parents=True)
            with patch.dict(os.environ,DAILY_ARTIFACT_MODULE=str(root),DAILY_NODE='node.exe'):
                self.assertEqual(portable_excel.artifact_runtime(),('node.exe',root))
                self.assertFalse(portable_excel.install_if_missing())
            manifest=json.loads((config.ROOT/'scripts/package.json').read_text())
            self.assertNotIn('@oai/artifact-tool',manifest.get('dependencies',{}))
            self.assertNotIn('@oai/artifact-tool',manifest.get('optionalDependencies',{}))


if __name__=='__main__':unittest.main()
