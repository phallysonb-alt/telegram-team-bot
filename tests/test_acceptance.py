import asyncio
import csv
import io
import json
import sqlite3
import unittest
import zipfile
from pathlib import Path
from datetime import datetime, timedelta
from types import SimpleNamespace
from team_bot.config import Config, ROOT
from team_bot.service import Service, safe_cell
from team_bot.reports import parse_report, parse_progress, NUMERIC
from team_bot.demo import seed, report, DEMO_DAY, EXPECTED
from team_bot.scheduler import Scheduler
from team_bot.knowledge import Knowledge
from team_bot.ai import AI
from team_bot.maintenance import backup, restore, InstanceLock


class Base(unittest.TestCase):
    def setUp(self):
        import uuid
        self.root = ROOT / 'data' / 'test-runs' / uuid.uuid4().hex
        self.c = Config(owner=99999, chat=-100123, data=self.root, backups=self.root / 'backups', report_times_confirmed=True, threads=dict(progress=11, training=22, management=33, report=44)).validate()
        self.now = datetime(2026, 10, 10, 20, 0, tzinfo=self.c.zone)
        self.c.now = lambda: self.now
        self.s = Service(self.c)

    def tearDown(self):
        self.s.db.close()
        # 测试数据保留在项目数据目录以便核对，不使用系统临时目录。

    def emp(self, i=1):
        self.s.register(10000 + i, f'E{i:03}', f'模拟员工{i:02}')
        self.s.approve(99999, 10000 + i)
        return 10000 + i


class ReportsTests(Base):
    def test_fixed_thirty_person_totals(self):
        r = seed(self.s)
        fixed = json.loads((ROOT / 'samples/固定数据期望.json').read_text(encoding='utf-8'))
        self.assertEqual(r['totals'], EXPECTED)
        self.assertEqual(r['totals'], fixed['全天已知'])
        self.assertEqual((r['expected'], r['submitted'], r['missing_reports'], r['leave']), (28, 21, 7, 1))
        self.assertEqual(r['missing_fields'], fixed['待补人数'])
        self.assertFalse(r['complete'])
        self.assertEqual(r['management']['completed'], 0)
        self.assertEqual(r['management']['overdue'], 1)

    def test_zero_missing_negative_and_format(self):
        good = report(DEMO_DAY, 1)
        _, p = parse_report(good.replace('上半场精聊新增：1', '上半场精聊新增：0'), self.now.date(), 'E001')
        self.assertEqual(p[NUMERIC[0]], 0)
        _, p = parse_report(good.replace('上半场精聊新增：1', '上半场精聊新增：'), self.now.date(), 'E001')
        self.assertIsNone(p[NUMERIC[0]])
        for value in ('-1', '1.5', '十', '+2', '1,000'):
            with self.assertRaises(ValueError):
                parse_report(good.replace('上半场精聊新增：1', '上半场精聊新增：' + value), self.now.date(), 'E001')
        with self.assertRaises(ValueError):
            parse_report('错误格式', self.now.date(), 'E001')

    def test_future_and_wrong_code(self):
        with self.assertRaises(ValueError):
            parse_report(report('2026-10-11', 1), self.now.date(), 'E001')
        with self.assertRaises(ValueError):
            parse_report(report(DEMO_DAY, 1), self.now.date(), 'E002')
        with self.assertRaises(ValueError):
            parse_report(report('2026-1-1', 1), self.now.date(), 'E001')

    def test_progress_and_final_share_one_row(self):
        uid = self.emp()
        self.s.progress(uid, f'日期：{DEMO_DAY}\n阶段：上半场\n精聊新增：100\n引流新增：100\n维聊新增：100', 'msg1')
        draft, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), 'msg2')
        self.s.confirm(uid, draft)
        r = self.s.summary(DEMO_DAY)
        self.assertEqual(r['totals'], dict(精聊=2, 引流=6, 维聊=8))
        self.assertEqual(self.s.db.one('SELECT COUNT(*) AS n FROM records')['n'], 1)
        with self.assertRaises(ValueError):
            self.s.progress(uid, f'日期：{DEMO_DAY}\n阶段：上半场\n精聊新增：999', 'msg3')

    def test_edit_preview_invalidates_old_and_repeat_confirm(self):
        uid = self.emp()
        first, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), 'source')
        second, _, _ = self.s.draft(uid, report(DEMO_DAY, 1).replace('上半场精聊新增：1', '上半场精聊新增：9'), 'source')
        with self.assertRaises(ValueError):
            self.s.confirm(uid, first)
        self.s.confirm(uid, second)
        v = self.s.record(uid, DEMO_DAY)['version']
        self.assertIn('未重复', self.s.confirm(uid, second))
        self.assertEqual(self.s.record(uid, DEMO_DAY)['version'], v)
        self.assertEqual(self.s.summary(DEMO_DAY)['totals']['精聊'], 10)

    def test_admin_correction_audit_revoke_late(self):
        uid = self.emp()
        d, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), '1')
        self.s.confirm(uid, d)
        self.assertEqual(self.s.record(uid, DEMO_DAY)['late'], 1)
        self.s.correct(99999, uid, DEMO_DAY, report(DEMO_DAY, 1).replace('上半场精聊新增：1', '上半场精聊新增：7'), '核实数量')
        audit = self.s.db.one('SELECT * FROM history ORDER BY id DESC LIMIT 1')
        self.assertEqual(audit['actor'], 99999)
        self.assertEqual(audit['reason'], '核实数量')
        self.assertIn('上半场精聊新增', audit['before_value'])
        self.s.revoke(uid, uid, DEMO_DAY, '填写错误')
        self.assertIsNone(self.s.record(uid, DEMO_DAY))
        self.assertEqual(self.s.summary(DEMO_DAY)['totals']['精聊'], 0)

    def test_employee_cannot_modify_another(self):
        uid, other = self.emp(), self.emp(2)
        d, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), '1')
        with self.assertRaises(PermissionError):
            self.s.confirm(other, d)
        with self.assertRaises(PermissionError):
            self.s.correct(other, uid, DEMO_DAY, report(DEMO_DAY, 1), 'test')
        with self.assertRaises(PermissionError):
            self.s.revoke(other, uid, DEMO_DAY, 'test')

    def test_deadline_and_open_window(self):
        uid = self.emp()
        self.now = self.now.replace(hour=17)
        with self.assertRaises(ValueError):
            self.s.draft(uid, report('2026-10-10', 1), '1')
        self.now = self.now.replace(hour=20)
        d, _, _ = self.s.draft(uid, report('2026-10-10', 1), '1')
        self.s.confirm(uid, d)
        self.assertEqual(self.s.record(uid, '2026-10-10')['late'], 0)

    def test_targets_not_invented(self):
        self.emp()
        self.assertEqual(self.s.summary(DEMO_DAY)['goals'], [])
        self.s.set_metric(99999, '精聊', '管理员定义', 10)
        self.assertEqual(self.s.summary(DEMO_DAY)['goals'][0]['gap'], 10)

    def test_csv_bom_and_formula_injection(self):
        uid = self.emp()
        text = report(DEMO_DAY, 1).replace('今日主要工作：完成培训和工作记录', '今日主要工作：=HYPERLINK("x")')
        d, _, _ = self.s.draft(uid, text, '1')
        self.s.confirm(uid, d)
        files = self.s.export_day(DEMO_DAY)
        path = next(f for f in files if f.suffix == '.csv')
        self.assertTrue(path.read_bytes().startswith(b'\xef\xbb\xbf'))
        self.assertIn("'=HYPERLINK", path.read_text(encoding='utf-8-sig'))
        self.assertEqual(safe_cell(' @test'), "' @test")

    def test_leave_and_disabled_missing_list(self):
        self.emp()
        self.emp(2)
        self.emp(3)
        self.s.leave(99999, 10002, DEMO_DAY, 1, '请假')
        self.s.change_employee(99999, 10003, active=0, reason='停用')
        self.assertEqual(self.s.missing(DEMO_DAY), ['E001'])
        self.assertEqual(self.s.missing(DEMO_DAY, '上半场'), ['E001'])

    def test_draft_survives_database_reopen(self):
        uid = self.emp()
        d, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), '1')
        self.s.db.close()
        self.s = Service(self.c)
        self.s.confirm(uid, d)
        self.assertEqual(self.s.summary(DEMO_DAY)['submitted'], 1)

    def test_partial_final_empty_overwrites_known_progress(self):
        uid = self.emp()
        self.s.progress(uid, f'日期：{DEMO_DAY}\n阶段：下半场\n精聊新增：2\n引流新增：100\n维聊新增：1', '1')
        d, _, _ = self.s.draft(uid, report(DEMO_DAY, 1, incomplete=True), '2')
        self.s.confirm(uid, d)
        r = self.s.summary(DEMO_DAY)
        self.assertEqual(r['totals']['引流'], 2)
        self.assertEqual(r['missing_fields']['下半场引流新增'], 1)

    def test_employee_csv_transaction_and_role_not_imported(self):
        self.s.import_employees(99999, '员工编号,姓名,用户ID,在职,需要日报\nA,甲,12345,1,1\n')
        self.assertFalse(self.s.admin(12345))
        with self.assertRaises(ValueError):
            self.s.import_employees(99999, '员工编号,姓名,用户ID,在职,需要日报\nB,乙,12346,1,1\nC,丙,-1,1,1\n')
        self.assertIsNone(self.s.db.one('SELECT * FROM employees WHERE uid=12346'))

    def test_today_roster_changes_and_missing_text(self):
        uid = self.emp()
        today = self.now.date().isoformat()
        self.s.roster(today)
        self.emp(2)
        self.assertEqual(len(self.s.roster(today)[0]), 2)
        self.s.change_employee(99999, 10002, active=0, reason='停用')
        self.assertEqual(self.s.roster(today)[0], [uid])
        d, _, _ = self.s.draft(uid, report(today, 1).replace('明日待办：无', '明日待办：'), '1')
        self.s.confirm(uid, d)
        self.assertFalse(self.s.summary(today)['complete'])
        self.assertEqual(self.s.summary(today)['missing_texts']['明日待办'], 1)

    def test_unknown_historical_roster_not_claimed_complete(self):
        uid = self.emp()
        d, _, _ = self.s.draft(uid, report(DEMO_DAY, 1), '1')
        self.s.confirm(uid, d)
        r = self.s.summary(DEMO_DAY)
        self.assertFalse(r['complete'])
        self.assertIn('尚未核实', r['roster_warning'])
        self.s.set_roster(99999, DEMO_DAY, [uid], '核对该日在职名单')
        self.assertTrue(self.s.summary(DEMO_DAY)['complete'])


class TasksAndSchedulerTests(Base):
    def test_multi_owner_completion_and_edit_audit(self):
        one, two = self.emp(), self.emp(2)
        task = self.s.new_task(99999, DEMO_DAY, '培训', '阅读', [one, two], DEMO_DAY + 'T21:00')
        self.s.task_update(one, task, '已完成', '完成')
        self.assertFalse(self.s.task(task)['all_done'])
        self.s.task_update(two, task, '已完成', '完成')
        self.assertTrue(self.s.task(task)['all_done'])
        self.s.task_edit(99999, task, deadline=DEMO_DAY + 'T22:00', reason='延期')
        self.assertEqual(self.s.db.one('SELECT reason FROM history ORDER BY id DESC LIMIT 1')['reason'], '延期')

    def test_admin_permission_all_operations(self):
        uid = self.emp()
        for action in [lambda: self.s.approve(uid, uid), lambda: self.s.change_employee(uid, uid, reason='x', role='admin'), lambda: self.s.new_task(uid, DEMO_DAY, 'x', 'x', [uid], DEMO_DAY + 'T20:00'), lambda: self.s.leave(uid, uid, DEMO_DAY, 1, 'x'), lambda: self.s.set_metric(uid, '精聊', 'x', 10), lambda: Knowledge(self.s).reindex(uid)]:
            with self.assertRaises(PermissionError):
                action()

    def activate(self, sch):
        sch.confirm(99999, self.c.timezone, ','.join(self.c.times.values()))

    def test_default_disabled_and_config_change_disables(self):
        sch = Scheduler(self.s)
        sch.plan(self.now)
        self.assertEqual(self.s.db.all('SELECT * FROM jobs'), [])
        self.activate(sch)
        self.assertTrue(sch.schedule_enabled())
        self.c.times['am'] = '11:00'
        self.assertFalse(sch.schedule_enabled())

    def test_restart_scheduler_once_and_topic_routing(self):
        self.emp()
        self.now = datetime(2026, 10, 9, 12, 1, tzinfo=self.c.zone)
        sch = Scheduler(self.s)
        self.activate(sch)
        sch.plan(self.now)
        first = self.s.db.all('SELECT * FROM outbox')
        self.assertEqual(len(first), 1)
        self.assertEqual(first[0]['thread'], 11)
        self.s.db.close()
        self.s = Service(self.c)
        sch = Scheduler(self.s)
        sch.plan(self.now)
        self.assertEqual(len(self.s.db.all('SELECT * FROM outbox')), 1)
        sch.publish(DEMO_DAY, 'manual-test')
        threads = {r['thread'] for r in self.s.db.all('SELECT * FROM outbox')}
        self.assertEqual(threads, {11, 33, 44})
        v = self.s.db.one('SELECT MAX(version) AS v FROM summaries')['v']
        sch.publish(DEMO_DAY, 'manual-test')
        self.assertEqual(self.s.db.one('SELECT MAX(version) AS v FROM summaries')['v'], v)
        sch.enqueue('training-test', 'training', '培训')
        self.assertEqual(self.s.db.one("SELECT thread FROM outbox WHERE key='training-test:0'")['thread'], 22)

    def test_midnight_catchup_limits_and_weekend(self):
        self.emp()
        self.now = datetime(2026, 10, 10, 0, 5, tzinfo=self.c.zone)
        sch = Scheduler(self.s)
        self.activate(sch)
        sch.plan(self.now)
        self.assertFalse(any(r['day'] == '2026-10-10' for r in self.s.db.all('SELECT * FROM jobs')))
        self.assertEqual(self.s.db.one("SELECT state FROM jobs WHERE key='2026-10-09:summary'")['state'], 'expired')
        self.assertEqual(self.s.db.all('SELECT * FROM outbox'), [])

    def test_cross_timezone_date(self):
        self.c.timezone = 'America/Los_Angeles'
        utc = datetime.fromisoformat('2026-10-10T02:00:00+00:00')
        local = utc.astimezone(self.c.zone)
        self.assertEqual(local.date().isoformat(), '2026-10-09')
        self.c.now = lambda: local
        sch = Scheduler(self.s)
        self.activate(sch)
        sch.plan(local)
        self.assertTrue(all(datetime.fromisoformat(j['due']).utcoffset() == local.utcoffset() for j in self.s.db.all('SELECT * FROM jobs')))

    def test_coalesced_refresh_versions(self):
        uid = self.emp()
        sch = Scheduler(self.s)
        self.activate(sch)
        sch.publish(DEMO_DAY, 'first')
        self.s.correct(99999, uid, DEMO_DAY, report(DEMO_DAY, 1), '补报')
        self.s.correct(99999, uid, DEMO_DAY, report(DEMO_DAY, 1).replace('上半场精聊新增：1', '上半场精聊新增：3'), '更正')
        self.assertEqual(len(self.s.db.all('SELECT * FROM dirty')), 1)
        sch.refresh(self.now)
        self.assertEqual(self.s.db.one('SELECT MAX(version) AS v FROM summaries')['v'], 1)
        self.now += timedelta(seconds=181)
        sch.refresh(self.now)
        self.assertEqual(self.s.db.one('SELECT MAX(version) AS v FROM summaries')['v'], 2)
        sch.refresh(self.now)
        self.assertEqual(self.s.db.one('SELECT MAX(version) AS v FROM summaries')['v'], 2)


class KnowledgeTests(Base):
    def test_chinese_sources_version_and_not_found(self):
        kb = Knowledge(self.s)
        p = self.root / '手册.md'
        p.write_text('# 日报口径\n上半场与下半场均填写本阶段新增量。', encoding='utf-8')
        first, _ = kb.import_file(99999, p)
        hits = kb.search('阶段新增量')
        self.assertTrue(hits)
        self.assertIn('日报口径', kb.local_answer(hits))
        p.write_text('# 新版口径\n只使用全新规则代号。', encoding='utf-8')
        second, version = kb.import_file(99999, p)
        self.assertEqual(version, 2)
        self.assertEqual(self.s.db.one('SELECT active FROM documents WHERE id=?', (first,))['active'], 0)
        self.assertEqual(kb.search('阶段新增量'), [])
        self.assertEqual(kb.local_answer(kb.search('毫不相关查询')), '未找到对应资料')
        kb.change(99999, second, '停用')
        self.assertEqual(kb.search('全新规则'), [])

    def test_scan_pdf_rejected_and_docx_parsed(self):
        from pypdf import PdfWriter
        from docx import Document
        blank = self.root / '扫描.pdf'
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        w.write(blank)
        with self.assertRaisesRegex(ValueError, '扫描 PDF'):
            Knowledge(self.s).import_file(99999, blank)
        p = self.root / '培训.docx'
        doc = Document()
        doc.add_heading('培训章节', 1)
        doc.add_paragraph('工作新增量应分阶段填写。')
        doc.save(p)
        Knowledge(self.s).import_file(99999, p)
        self.assertTrue(Knowledge(self.s).search('阶段填写'))

    def test_fallback_and_delete(self):
        self.s.db.fts = False
        p = self.root / '资料.txt'
        p.write_text('团队培训资料规范。', encoding='utf-8')
        kb = Knowledge(self.s)
        d, _ = kb.import_file(99999, p)
        self.assertTrue(kb.search('培训资料'))
        kb.change(99999, d, '删除')
        self.assertEqual(kb.search('培训资料'), [])
        self.assertEqual(self.s.db.all('SELECT * FROM chunks'), [])

    def test_size_type_and_failed_new_version(self):
        p = self.root / '文件.exe'
        p.write_text('abc')
        with self.assertRaises(ValueError):
            Knowledge(self.s).import_file(99999, p)
        p = self.root / '资料.txt'
        p.write_text('原始有效培训资料', encoding='utf-8')
        Knowledge(self.s).import_file(99999, p)
        p.write_text('', encoding='utf-8')
        with self.assertRaises(ValueError):
            Knowledge(self.s).import_file(99999, p)
        self.assertTrue(Knowledge(self.s).search('原始有效'))


class AITests(Base, unittest.IsolatedAsyncioTestCase):
    def enable(self):
        self.c.ai_enabled = self.c.free_confirmed = True
        self.c.key, self.c.model = 'test-key', 'test-model'

    async def test_missing_config_fallback_accurate_totals(self):
        r = seed(self.s)
        result, status = await AI(self.s).generate('汇总', '培训说明')
        self.assertIsNone(result)
        self.assertEqual(r['totals'], EXPECTED)
        self.assertIn('未使用 AI', self.s.render_summary(r))

    async def test_timeout_and_429_count_each_retry(self):
        self.enable()
        async def timeout(text):
            raise asyncio.TimeoutError()
        result, _ = await AI(self.s, timeout).generate('测试', '一般培训说明')
        self.assertIsNone(result)
        self.assertEqual(self.s.db.one("SELECT COUNT(*) AS n FROM ai_calls WHERE status='失败'")['n'], 1)
        self.c.retries = 1
        async def limit(text):
            raise RuntimeError('429')
        result, _ = await AI(self.s, limit).generate('测试', '一般培训说明')
        self.assertIsNone(result)
        self.assertEqual(self.s.db.one("SELECT COUNT(*) AS n FROM ai_calls WHERE status='限流'")['n'], 2)
        self.assertEqual(self.s.summary(DEMO_DAY)['totals'], dict(精聊=0, 引流=0, 维聊=0))

    async def test_budget_cache_and_sensitive_screen(self):
        self.enable()
        self.emp()
        async def good(text):
            return '根据培训资料填写。'
        ai = AI(self.s, good)
        self.c.budget = 1
        result, _ = await ai.generate('培训', '阶段填写说明')
        self.assertIsNotNone(result)
        _, status = await ai.generate('培训', '阶段填写说明')
        self.assertEqual(status, 'AI 缓存')
        _, status = await ai.generate('培训', '另一条培训说明')
        self.assertIn('预算', status)
        for text in ('电话：123456789', '客户姓名：某人', '模拟员工01遇到问题', '邮箱 a@b.com', '0xdeadbeef1234567890'):
            self.assertFalse(ai.safe(text))

    async def test_confidential_docs_never_leave(self):
        self.enable()
        calls = []
        async def capture(text):
            calls.append(text)
            return '回答'
        p = self.root / '内部.txt'
        p.write_text('内部培训规则：待办优先记录。', encoding='utf-8')
        kb = Knowledge(self.s)
        d, _ = kb.import_file(99999, p)
        hits = kb.search('待办记录')
        ai = AI(self.s, capture)
        answer = await ai.ask('待办记录', hits)
        self.assertIn('来源', answer)
        self.assertEqual(calls, [])
        kb.change(99999, d, '外发允许')
        await ai.ask('待办记录', kb.search('待办记录'))
        self.assertEqual(len(calls), 1)
        await ai.ask('电话 123456789 的待办记录', kb.search('待办记录'))
        self.assertEqual(len(calls), 1)

    async def test_summary_ai_only_approved_text_once_and_numbers_unchanged(self):
        self.enable()
        seed(self.s)
        calls = []
        async def capture(text):
            calls.append(text)
            return '培训已经开展，待办已整理。'
        self.s.db.set('public_notes:' + DEMO_DAY, '培训已经开展，待办已整理。')
        sch = Scheduler(self.s)
        sch.ai = AI(self.s, capture)
        result = sch.publish(DEMO_DAY, 'summary-ai-test')
        await sch.enhance()
        await sch.enhance()
        self.assertEqual(len(calls), 1)
        self.assertNotIn('10001', calls[0])
        self.assertNotIn('模拟员工', calls[0])
        saved = json.loads(self.s.db.one('SELECT payload FROM summaries WHERE day=? AND version=?', (DEMO_DAY, result['version']))['payload'])
        self.assertEqual(saved['totals'], EXPECTED)
        self.assertIn('培训', saved['ai_text'])


class ReliabilityTests(Base):
    def test_backup_restore_history(self):
        seed(self.s)
        kb = Knowledge(self.s)
        p = self.root / '手册.md'
        p.write_text('培训资料内容。', encoding='utf-8')
        kb.import_file(99999, p)
        count = self.s.db.one('SELECT COUNT(*) AS n FROM history')['n']
        archive = backup(self.s)
        self.s.db.close()
        target = self.root / 'restored'
        other = Config(owner=99999, data=target, backups=target / 'backups').validate()
        restore(other, archive)
        self.s = Service(other)
        self.assertEqual(self.s.db.one('SELECT COUNT(*) AS n FROM history')['n'], count)
        self.assertEqual(len(self.s.db.all('SELECT * FROM employees')), 30)
        self.assertTrue(Knowledge(self.s).search('培训资料'))
        self.assertEqual(self.s.db.run('PRAGMA integrity_check').fetchone()[0], 'ok')

    def test_lock_and_unsafe_restore(self):
        with InstanceLock(self.root):
            with self.assertRaises(RuntimeError):
                with InstanceLock(self.root):
                    pass
        archive = self.root / 'unsafe.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('team.sqlite3', b'bad')
            z.writestr('../escape.txt', b'bad')
        with self.assertRaises(ValueError):
            restore(self.c, archive)


class TelegramAdapterTests(Base, unittest.IsolatedAsyncioTestCase):
    def make_update(self, update_id, text='', thread=44, uid=10001, chat=None, callback=None, edited=False):
        from telegram import Update, Bot
        bot = Bot('123:fake')
        payload = dict(update_id=update_id)
        msg = dict(message_id=50, date=1791590400, chat=chat or dict(id=self.c.chat, type='supergroup', title='模拟群', is_forum=True), from_user=dict(id=uid, first_name='模拟', is_bot=False), message_thread_id=thread, text=text)
        msg['from'] = msg.pop('from_user')
        if callback:
            payload['callback_query'] = dict(id='q', **{'from': msg['from']}, chat_instance='x', message=msg, data=callback)
        else:
            payload['edited_message' if edited else 'message'] = msg
        return Update.de_json(payload, bot)

    async def test_duplicate_update_atomic_and_edited_draft(self):
        from team_bot.bot import Adapter
        self.emp()
        adapter = Adapter(self.s)
        ctx = SimpleNamespace(bot=SimpleNamespace(id=123))
        u = self.make_update(1, report(DEMO_DAY, 1))
        await adapter.handle(u, ctx)
        await adapter.handle(u, ctx)
        self.assertEqual(self.s.db.one('SELECT COUNT(*) AS n FROM drafts')['n'], 1)
        await adapter.handle(self.make_update(2, report(DEMO_DAY, 1).replace('上半场精聊新增：1', '上半场精聊新增：2'), edited=True), ctx)
        self.assertEqual(self.s.db.one("SELECT COUNT(*) AS n FROM drafts WHERE state='draft'")['n'], 1)
        self.assertEqual(self.s.db.one('SELECT COUNT(*) AS n FROM records')['n'], 0)

    async def test_wrong_topic_admin_denied_and_ordinary_chat_no_ai(self):
        from team_bot.bot import Adapter
        self.emp()
        adapter = Adapter(self.s)
        ctx = SimpleNamespace(bot=SimpleNamespace(id=123))
        await adapter.handle(self.make_update(1, report(DEMO_DAY, 1), thread=11), ctx)
        self.assertEqual(self.s.db.all('SELECT * FROM drafts'), [])
        self.assertIn('对应话题', self.s.db.one('SELECT text FROM outbox')['text'])
        await adapter.handle(self.make_update(2, '/approve 10001'), ctx)
        self.assertTrue(any('仅管理员' in r['text'] for r in self.s.db.all('SELECT * FROM outbox')))
        await adapter.handle(self.make_update(3, '普通聊天'), ctx)
        self.assertEqual(self.s.db.all('SELECT * FROM ai_calls'), [])

    async def test_callback_confirmation_and_edit_does_not_silently_change(self):
        from unittest.mock import patch, AsyncMock
        from telegram import CallbackQuery
        from team_bot.bot import Adapter
        self.emp()
        adapter = Adapter(self.s)
        ctx = SimpleNamespace(bot=SimpleNamespace(id=123))
        await adapter.handle(self.make_update(1, report(DEMO_DAY, 1)), ctx)
        draft_id = self.s.db.one('SELECT id FROM drafts')['id']
        with patch.object(CallbackQuery, 'answer', new=AsyncMock()):
            await adapter.handle(self.make_update(2, callback=f'confirm:{draft_id}'), ctx)
            await adapter.handle(self.make_update(2, callback=f'confirm:{draft_id}'), ctx)
            await adapter.handle(self.make_update(3, callback=f'confirm:{draft_id}'), ctx)
        row = self.s.record(10001, DEMO_DAY)
        self.assertEqual((row['confirmed'], row['version']), (1, 1))
        await adapter.handle(self.make_update(4, report(DEMO_DAY, 1).replace('上半场精聊新增：1', '上半场精聊新增：99'), edited=True), ctx)
        self.assertEqual(self.s.summary(DEMO_DAY)['totals']['精聊'], 2)
        self.assertEqual(self.s.record(10001, DEMO_DAY)['version'], 1)

    async def test_binding_owner_and_unknown_group(self):
        from team_bot.bot import Adapter
        adapter = Adapter(self.s)
        ctx = SimpleNamespace(bot=SimpleNamespace(id=123))
        await adapter.handle(self.make_update(1, '/bind 培训', thread=222, uid=99999), ctx)
        self.assertEqual(adapter.scheduler.thread('training'), 222)
        await adapter.handle(self.make_update(2, '/bind 日报', thread=222, uid=99999), ctx)
        self.assertEqual(adapter.scheduler.thread('report'), 44)
        foreign = dict(id=-100999, type='supergroup', title='其他群', is_forum=True)
        await adapter.handle(self.make_update(3, '/bind 进度', thread=777, uid=99999, chat=foreign), ctx)
        self.assertEqual(adapter.scheduler.thread('progress'), 11)

    async def test_split_delivery_and_network_recovery(self):
        sch = Scheduler(self.s)
        sch.enqueue('long', 'training', '中' * 8000)
        self.assertEqual(len(self.s.db.all('SELECT * FROM outbox')), 3)
        class FakeBot:
            def __init__(self):
                self.fail = True
                self.sent = []
            async def send_message(bot, **kwargs):
                if bot.fail:
                    raise TimeoutError()
                bot.sent.append(kwargs)
                return SimpleNamespace(message_id=len(bot.sent))
        bot = FakeBot()
        await sch.flush(bot)
        self.assertTrue(all(r['attempts'] == 1 and r['state'] == 'pending' for r in self.s.db.all('SELECT * FROM outbox')))
        bot.fail = False
        self.now += timedelta(seconds=61)
        await sch.flush(bot)
        self.assertEqual(len(bot.sent), 3)
        self.assertTrue(all(r['message_thread_id'] == 22 for r in bot.sent))
        await sch.flush(bot)
        self.assertEqual(len(bot.sent), 3)

    def test_application_and_sdk_construct_offline(self):
        from telegram.ext import Application
        from google import genai
        from google.genai import types
        app = Application.builder().token('123:fake').concurrent_updates(False).build()
        self.assertEqual(app.bot.token, '123:fake')
        client = genai.Client(api_key='fake', http_options=types.HttpOptions(timeout=2000, retry_options=types.HttpRetryOptions(attempts=1)))
        client.close()
