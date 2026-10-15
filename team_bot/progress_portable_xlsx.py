"""Template-preserving fallback when artifact-tool is absent on another company PC."""
from copy import copy
from datetime import date
import math

def write(data,dest):
 from openpyxl import Workbook,load_workbook
 from openpyxl.utils import get_column_letter
 from .config import ROOT
 template=load_workbook(ROOT/'templates/Telegram_Progress_History_Template_v2.xlsx');wb=Workbook();wb.remove(wb.active)
 def lines(value,width):return sum(max(1,math.ceil(sum(2 if ord(c)>255 else 1 for c in line)/max(7,width-1))) for line in str(value or '').split('\n'))
 def cell(sh,row,col,value,source):
  if isinstance(value,str) and len(value)>32767:raise ValueError('文本超过Excel单元格上限；原记录不变')
  c=sh.cell(row,col,value)
  for field in ('font','fill','border','alignment','protection'):setattr(c,field,copy(getattr(source,field)))
  c.number_format=source.number_format
  if isinstance(value,str):c.data_type='s'
  if isinstance(value,date):c.number_format='yyyy-mm-dd'
  return c
 for name,last,source_head,records,date_col in [('查询汇总',12,5,data['summaries'],0),('员工进度历史',19,6,None,0),('当前问题历史',12,5,data['problems'],0),('任务定义',10,5,data['definitions'],1)]:
  ref=template[name];sh=wb.create_sheet(name);sh.sheet_view.showGridLines=False;sh.freeze_panes='A6'
  widths=[]
  for col in range(1,last+1):
   key=get_column_letter(col);width=ref.column_dimensions[key].width;widths.append(width);sh.column_dimensions[key].width=width
  def merged(row,text,source_row,height):
   for col in range(1,last+1):cell(sh,row,col,None,ref.cell(source_row,col))
   sh.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last);cell(sh,row,1,text,ref.cell(source_row,1));sh.row_dimensions[row].height=height
  for row in range(1,4):merged(row,data['label'] if row==3 else str(ref.cell(row,1).value).replace('（修正版式样板）',''),row,28 if row==1 else 24)
  def header(row):
   for col in range(1,last+1):cell(sh,row,col,str(ref.cell(source_head,col).value).replace('身份ID（演示）','身份ID'),ref.cell(source_head,col))
   sh.row_dimensions[row].height=ref.row_dimensions[source_head].height
  def rows(row,values):
   for original in values:
    for index,value in enumerate(original):
     if index==date_col and value:value=date.fromisoformat(value)
     source=ref.cell(source_head+1,1)
     if name=='员工进度历史' and value=='已完成' and (5<=index<=8 or index==18):source=template['员工进度历史']['F7']
     elif name=='员工进度历史' and value=='未反馈' and index==4:source=template['员工进度历史']['E10']
     elif name=='员工进度历史' and value=='未反馈' and 5<=index<=8:source=template['员工进度历史']['F10']
     c=cell(sh,row,index+1,value,source)
     if name=='查询汇总' and index==9:c.number_format='0%'
    sh.row_dimensions[row].height=min(180,max(ref.row_dimensions[source_head+1].height or 24,*[16*lines(v,widths[i])+4 for i,v in enumerate(original)]));row+=1
   return row
  if records is not None:header(5);rows(6,records)
  else:
   row=5
   for group in data['groups']:
    t=group['task'];text=f"任务：{t['title']}    ｜    规则：{'按顺序完成' if t['kind']=='seq' else '自由完成'}    ｜    阶段总数：{len(t['stages'])}";merged(row,text,5,max(24,16*lines(text,sum(widths))+4));header(row+1);row=rows(row+2,group['rows'])+1
 wb.save(dest)
