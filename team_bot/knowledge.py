import re
import hashlib
from pathlib import Path
from .db import stamp


def tokens(text):
    """中文双字切片 + 单字回退，英文数字词；不用联网分词。"""
    groups = re.findall(r'[\u4e00-\u9fff]+|[a-zA-Z0-9]+', text.lower())
    result = []
    for group in groups:
        if '\u4e00' <= group[0] <= '\u9fff':
            result.extend(group[i:i + 2] for i in range(len(group) - 1))
            if len(group) == 1:
                result.append(group)
        else:
            result.append(group)
    return list(dict.fromkeys(result))


def extract(path):
    suffix = path.suffix.lower()
    if suffix in ('.txt', '.md'):
        text = path.read_text(encoding='utf-8-sig')
        sections, title, buf = [], '正文', []
        for line in text.splitlines():
            if line.startswith('#') and buf:
                sections.append((title, '\n'.join(buf)))
                buf = []
            if line.startswith('#'):
                title = line.lstrip('#').strip() or '正文'
            else:
                buf.append(line)
        sections.append((title, '\n'.join(buf)))
        return sections
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        sections, title, buf = [], '正文', []
        for p in document.paragraphs:
            if p.style and p.style.name.startswith('Heading'):
                sections.append((title, '\n'.join(buf)))
                title, buf = p.text, []
            else:
                buf.append(p.text)
        for table in document.tables:
            buf.extend(' | '.join(c.text for c in r.cells) for r in table.rows)
        sections.append((title, '\n'.join(buf)))
        return sections
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('不支持加密 PDF')
        sections = [(f'第 {i + 1} 页', p.extract_text() or '') for i, p in enumerate(reader.pages)]
        if not any(text.strip() for _, text in sections):
            raise ValueError('扫描 PDF 暂不支持文字提取，请上传文字型 PDF 或 TXT')
        return sections
    raise ValueError('只支持 TXT、Markdown、DOCX、文字型 PDF')


class Knowledge:
    def __init__(self, service):
        self.s = service
        self.db = service.db

    def import_file(self, actor, path):
        self.s.require_admin(actor)
        if self.s.c.training_mode=='directory':
            raise ValueError('目录模式禁止正文提取和索引，请使用资料目录入口')
        path = Path(path)
        if path.stat().st_size > self.s.c.max_file_mb * 1024 * 1024:
            raise ValueError('资料超过文件大小限制')
        sections = extract(path)
        pieces = [(loc, text[start:start + 1200]) for loc, text in sections for start in range(0, len(text), 1100) if text[start:start + 1200].strip()]
        if not pieces:
            raise ValueError('文件无可检索文字')
        # 解析成功后再停用旧版；原文件持久保存，文件名不作为路径。
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        dest = self.s.c.data / 'materials' / (digest + path.suffix.lower())
        if path.resolve() != dest.resolve():
            dest.write_bytes(path.read_bytes())
        with self.db.transaction():
            name = path.name
            version = self.db.one('SELECT COALESCE(MAX(version),0)+1 AS v FROM documents WHERE name=?', (name,))['v']
            old = self.db.all('SELECT * FROM documents WHERE name=? AND active=1', (name,))
            self.db.run('UPDATE documents SET active=0 WHERE name=?', (name,))
            doc = self.db.run('INSERT INTO documents(name,version,imported,path) VALUES (?,?,?,?)', (name, version, stamp(), str(dest))).lastrowid
            for location, text in pieces:
                chunk = self.db.run('INSERT INTO chunks(document,location,text) VALUES (?,?,?)', (doc, location, text)).lastrowid
                if self.db.fts:
                    self.db.run('INSERT INTO search_index(tokens,chunk) VALUES (?,?)', (' '.join(tokens(text)), chunk))
            self.db.audit('资料', doc, actor, '导入新版本，默认仅本地使用', old, dict(name=name, version=version))
        return doc, version

    def change(self, actor, doc, action):
        self.s.require_admin(actor)
        if self.s.c.training_mode=='directory' and action=='外发允许':
            raise ValueError('目录模式禁止资料正文外发模型')
        if action not in ('外发允许', '外发禁止', '停用', '删除'):
            raise ValueError('操作应为外发允许、外发禁止、停用或删除')
        with self.db.transaction():
            old = self.db.one('SELECT * FROM documents WHERE id=?', (doc,))
            if not old:
                raise ValueError('资料不存在')
            if action == '删除':
                if self.db.fts:
                    self.db.run('DELETE FROM search_index WHERE chunk IN (SELECT id FROM chunks WHERE document=?)', (doc,))
                self.db.run('DELETE FROM documents WHERE id=?', (doc,))
            elif action == '停用':
                self.db.run('UPDATE documents SET active=0 WHERE id=?', (doc,))
            else:
                self.db.run('UPDATE documents SET external=? WHERE id=?', (int(action == '外发允许'), doc))
            self.db.audit('资料', doc, actor, action, old, self.db.one('SELECT * FROM documents WHERE id=?', (doc,)))
        if action == '删除':
            path = Path(old['path']).resolve()
            materials = (self.s.c.data / 'materials').resolve()
            if path.parent == materials and not self.db.one('SELECT id FROM documents WHERE path=?', (str(path),)):
                path.unlink(missing_ok=True)

    def reindex(self, actor):
        self.s.require_admin(actor)
        if self.s.c.training_mode=='directory':
            raise ValueError('目录模式不建立正文索引')
        if not self.db.fts:
            return 'FTS5 不可用，已使用本地中文关键词回退'
        with self.db.transaction():
            self.db.run('DELETE FROM search_index')
            for r in self.db.all('SELECT id,text FROM chunks'):
                self.db.run('INSERT INTO search_index(tokens,chunk) VALUES (?,?)', (' '.join(tokens(r['text'])), r['id']))
            self.db.audit('资料', '索引', actor, '重建中文检索索引', None, '已重建')
        return '索引已重建'

    def search(self, query, limit=4):
        words = tokens(query)[:30]
        if not words:
            return []
        query_sql = 'SELECT c.*,d.name,d.version,d.external,d.imported FROM chunks c JOIN documents d ON d.id=c.document WHERE d.active=1'
        if self.db.fts:
            ids = [r['chunk'] for r in self.db.all('SELECT chunk FROM search_index WHERE search_index MATCH ? ORDER BY rank LIMIT 100', (' OR '.join('"' + w + '"' for w in words),))]
            if not ids:
                return []
            rows = self.db.all(query_sql + ' AND c.id IN (' + ','.join('?' for _ in ids) + ')', ids)
        else:
            rows = self.db.all(query_sql)
        scored = [(sum(w in r['text'].lower() for w in words), r) for r in rows]
        return [r for score, r in sorted(scored, key=lambda pair: pair[0], reverse=True) if score][:limit]

    @staticmethod
    def local_answer(hits):
        if not hits:
            return '未找到对应资料'
        return '\n\n'.join(f'来源：{h["name"]} · v{h["version"]} · {h["location"]}\n导入时间：{h["imported"]}\n{h["text"]}' for h in hits)
