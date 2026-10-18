"""Optional independent-company export provider; frozen daily logic is unchanged."""
import os,uuid,zipfile,math,shutil
from pathlib import Path
from datetime import date
def artifact_runtime():
 from .config import ROOT
 bundled=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
 paths=[bundled,ROOT/'scripts']
 if os.getenv('DAILY_ARTIFACT_MODULE'):paths=[Path(os.environ['DAILY_ARTIFACT_MODULE'])]
 node=os.getenv('DAILY_NODE') or shutil.which('node')
 if not node and (bundled/'bin/node.exe').is_file():node=str(bundled/'bin/node.exe')
 module=next((p for p in paths if (p/'node_modules/@oai/artifact-tool').exists()),None)
 return (node,module) if node and module else None

def available_artifact():return artifact_runtime() is not None

def install_if_missing():
 # Retain entrypoint API; select per call instead of monkeypatching globally.
 return not available_artifact()
def write_daily(data,destination):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Border,Side,Alignment
 from openpyxl.utils import get_column_letter
 path=Path(destination);path.parent.mkdir(parents=True,exist_ok=True);temp=path.parent/(uuid.uuid4().hex+'.xlsx');wb=Workbook();wb.remove(wb.active);blue='16324F';edge=Side(style='thin',color='D9D9D9')
 def init(name,widths,freeze):
  sh=wb.create_sheet(name);sh.sheet_view.showGridLines=False;sh.sheet_properties.tabColor=blue;sh.freeze_panes='A'+str(freeze+1)
  for i,width in enumerate(widths,1):sh.column_dimensions[get_column_letter(i)].width=width
  for row,value in [(1,'日报历史记录 · '+name),(2,data.get('label') or '数据库生成 · 只读历史报表')]:
   sh.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8);c=sh.cell(row,1,value);c.font=Font(name='Noto Sans CJK SC',size=17 if row==1 else 11,bold=row==1,color=blue);sh.row_dimensions[row].height=32 if row==1 else 26
  return sh
 def row(sh,n,values,head=False):
  for i,value in enumerate(values,1):
   c=sh.cell(n,i,value)
   if isinstance(value,str):c.data_type='s'
   c.font=Font(name='Noto Sans CJK SC',size=11,bold=head,color='FFFFFF' if head else blue);c.alignment=Alignment(horizontal='center' if head or sh.title=='每日汇总' or i not in (1,2,7) else 'left',vertical='center',wrap_text=not head);c.border=Border(left=edge,right=edge)
   if head:c.fill=PatternFill('solid',fgColor=blue)
   elif n%2:c.fill=PatternFill('solid',fgColor='F2F6FA')
   if isinstance(value,date):c.number_format='yyyy-mm-dd'
   elif isinstance(value,(int,float)):c.number_format='#,##0'
  sh.row_dimensions[n].height=26
 summary=init('每日汇总',[17,13,13,13,13,11,11,18],4);detail=init('员工明细',[20,28,11,11,11,11,46,20],2)
 row(summary,4,['日期','新增','维聊','有效','精聊','已交人数','应交人数','状态'],True);n=5;d=4
 for day in data['days']:
  row(summary,n,[date.fromisoformat(day['day']),*[day['totals'][k] if day['submitted'] else None for k in ('新增','维聊','有效','精聊')],day['submitted'],day['expected'],day['status']]);n+=1
  detail.merge_cells(start_row=d,start_column=1,end_row=d,end_column=8);c=detail.cell(d,1,date.fromisoformat(day['day']));c.number_format='yyyy-mm-dd';c.alignment=Alignment(horizontal='center',vertical='center');c.font=Font(name='Noto Sans CJK SC',size=15,bold=True,color=blue);c.fill=PatternFill('solid',fgColor='DEEAF5');detail.row_dimensions[d].height=30;d+=1
  row(detail,d,['员工名称','Telegram 用户名','新增','维聊','有效','精聊','备注','提交状态'],True);d+=1
  for v in day['rows']:
   remaining=str(v[6]);chunks=[]
   while remaining:
    end=min(250,len(remaining));breaks=[i for i,c in enumerate(remaining[:end]) if c=='\n']
    if len(breaks)>9:end=breaks[8]+1
    elif end<len(remaining):
     newline=remaining.rfind('\n',0,end+1)
     if newline>100:end=newline+1
    chunks.append(remaining[:end]);remaining=remaining[end:]
   for index,note in enumerate(chunks or ['']):
    cells=[*v[:6],note,v[7]] if not index else [v[0]+'（备注续）',v[1],None,None,None,None,note,'备注续'];row(detail,d,cells)
    detail.row_dimensions[d].height=min(409,max(27,18*math.ceil(len(cells[0])/9),18*math.ceil(len(str(cells[1]))/22),18*math.ceil(len(str(cells[7]))/10),18*sum(max(1,math.ceil(len(line)/21)) for line in note.split('\n'))));d+=1
  d+=1
 try:
  wb.save(temp)
  with zipfile.ZipFile(temp) as z:
   if z.testzip():raise ValueError('Excel完整性检查失败')
  os.replace(temp,path)
 finally:temp.unlink(missing_ok=True)
